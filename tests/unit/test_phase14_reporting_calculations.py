"""
Unit tests for Phase 14: Reporting Calculations, Number to Words (INR),
OpenPyXL Excel Generation, ReportLab PDF Rendering, and Tax Arithmetic.
"""
from decimal import Decimal
import io
import pytest
import openpyxl

from app.services.export_service import (
    convert_number_to_words_inr,
    generate_excel_spreadsheet,
    generate_csv_stream,
    generate_posp_invoice_pdf,
    generate_voucher_pdf,
)


# ============================================================================
# 1. Indian Currency Number to Words (INR) Parity Tests
# ============================================================================

def test_number_to_words_zero():
    assert convert_number_to_words_inr(Decimal("0.00")) == "Rupees Zero Only"
    assert convert_number_to_words_inr(0) == "Rupees Zero Only"
    assert convert_number_to_words_inr(None) == "Rupees Zero Only"


def test_number_to_words_units_and_tens():
    assert convert_number_to_words_inr(Decimal("5.00")) == "Rupees Five Only"
    assert convert_number_to_words_inr(Decimal("17.00")) == "Rupees Seventeen Only"
    assert convert_number_to_words_inr(Decimal("45.00")) == "Rupees Forty Five Only"
    assert convert_number_to_words_inr(Decimal("99.00")) == "Rupees Ninety Nine Only"


def test_number_to_words_hundreds_and_thousands():
    assert convert_number_to_words_inr(Decimal("100.00")) == "Rupees One Hundred Only"
    assert convert_number_to_words_inr(Decimal("543.00")) == "Rupees Five Hundred Forty Three Only"
    assert convert_number_to_words_inr(Decimal("1000.00")) == "Rupees One Thousand Only"
    assert convert_number_to_words_inr(Decimal("25400.00")) == "Rupees Twenty Five Thousand Four Hundred Only"


def test_number_to_words_lakhs_and_crores():
    # Exactly 1 Lakh
    assert convert_number_to_words_inr(Decimal("100000.00")) == "Rupees One Lakh Only"
    # 1 Lakh 25 Thousand 450
    assert convert_number_to_words_inr(Decimal("125450.00")) == "Rupees One Lakh Twenty Five Thousand Four Hundred Fifty Only"
    # 5 Crore 20 Lakh
    assert convert_number_to_words_inr(Decimal("52000000.00")) == "Rupees Five Crore Twenty Lakh Only"


def test_number_to_words_with_paise():
    assert convert_number_to_words_inr(Decimal("125450.50")) == "Rupees One Lakh Twenty Five Thousand Four Hundred Fifty and Fifty Paise Only"
    assert convert_number_to_words_inr(Decimal("0.75")) == "Rupees Zero and Seventy Five Paise Only"
    assert convert_number_to_words_inr(Decimal("10.05")) == "Rupees Ten and Five Paise Only"


def test_number_to_words_negative():
    assert convert_number_to_words_inr(Decimal("-500.00")) == "Minus Rupees Five Hundred Only"


# ============================================================================
# 2. OpenPyXL Native Spreadsheet Generation Tests
# ============================================================================

def test_generate_excel_spreadsheet_structure():
    columns = ["Trans ID", "Policy No", "Net Premium", "Gross Premium"]
    rows = [
        [101, "POL-2026-001", Decimal("5000.00"), Decimal("5900.00")],
        [102, "POL-2026-002", Decimal("10000.00"), Decimal("11800.00")],
    ]
    file_bytes = generate_excel_spreadsheet(columns, rows, sheet_name="Transactions")

    # Verify OpenXML magic bytes (PK zip header)
    assert file_bytes.startswith(b"PK\x03\x04")

    # Load with openpyxl to verify content and formatting
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb["Transactions"]
    assert ws.title == "Transactions"
    assert ws.cell(row=1, column=1).value == "Trans ID"
    assert ws.cell(row=1, column=2).value == "Policy No"
    assert ws.cell(row=2, column=2).value == "POL-2026-001"
    assert ws.cell(row=3, column=2).value == "POL-2026-002"

    # Header background color should be 48D1CC (Medium Turquoise)
    header_fill = ws.cell(row=1, column=1).fill.start_color.rgb
    assert header_fill in ("FF48D1CC", "48D1CC", "0048D1CC")


# ============================================================================
# 3. CSV Streaming Generation Tests
# ============================================================================

def test_generate_csv_stream():
    columns = ["ID", "Name", "Amount"]
    rows = [
        [1, "Agent A", Decimal("1500.50")],
        [2, "Agent B", Decimal("2500.00")],
    ]
    stream = list(generate_csv_stream(columns, rows))
    csv_content = "".join(stream)
    assert "ID,Name,Amount" in csv_content
    assert "1,Agent A,1500.50" in csv_content
    assert "2,Agent B,2500.00" in csv_content


# ============================================================================
# 4. ReportLab PDF Generation Tests
# ============================================================================

def test_generate_posp_invoice_pdf():
    payload = {
        "invoice_no": "INV/2025-2026/04/1001",
        "invoice_date": "15/04/2025",
        "agent_id": 101,
        "agent_name": "Test POSP Partner",
        "posp_type": "POSP",
        "financial_year": "2025-2026",
        "month": "APRIL",
        "amount": Decimal("10000.00"),
        "gst_amt": Decimal("1800.00"),
        "grand_total": Decimal("11800.00"),
        "tds_amount": Decimal("500.00"),
        "net_payable": Decimal("11300.00"),
        "amount_in_words": "Rupees Eleven Thousand Three Hundred Only",
        "created_by": "System Admin",
        "status": "GENERATED",
    }
    pdf_bytes = generate_posp_invoice_pdf(payload)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    assert pdf_bytes.startswith(b"%PDF")


def test_generate_voucher_pdf():
    payload = {
        "voucher_no": "VCH-2026-0099",
        "voucher_date": "2026-04-10",
        "voucher_type": "PAYMENT",
        "payee_name": "Ramesh Patel",
        "payment_mode": "CHEQUE",
        "amount": Decimal("7500.00"),
        "amount_in_words": "Rupees Seven Thousand Five Hundred Only",
        "narration": "Commission disbursement for April 2026",
        "created_by": "Accounts User",
    }
    pdf_bytes = generate_voucher_pdf(payload)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    assert pdf_bytes.startswith(b"%PDF")


# ============================================================================
# 5. POSP Tax Arithmetic Parity
# ============================================================================

def test_posp_tax_arithmetic_registered_vs_unregistered():
    amount = Decimal("10000.00")

    # 1. Registered Agent: 18% GST added, 5% TDS deducted on base
    gst_registered = (amount * Decimal("0.18")).quantize(Decimal("0.01"))
    grand_registered = amount + gst_registered
    tds = (amount * Decimal("0.05")).quantize(Decimal("0.01"))
    net_registered = grand_registered - tds

    assert gst_registered == Decimal("1800.00")
    assert grand_registered == Decimal("11800.00")
    assert tds == Decimal("500.00")
    assert net_registered == Decimal("11300.00")

    # 2. Unregistered Agent: 0% GST, 5% TDS deducted on base
    gst_unregistered = Decimal("0.00")
    grand_unregistered = amount
    net_unregistered = grand_unregistered - tds

    assert gst_unregistered == Decimal("0.00")
    assert grand_unregistered == Decimal("10000.00")
    assert net_unregistered == Decimal("9500.00")

"""
Phase 14 End-to-End Golden Parity & Lifecycle Verification Test.
Executes the full pipeline:
Policy Booking -> Commission Accounting -> POSP Invoice Generation -> TDS Register
-> Payment Advice -> Dashboard Matrix -> OpenXML / CSV Export.
"""
from datetime import datetime, date, timedelta
from decimal import Decimal
import io
import pytest
from httpx import AsyncClient
import openpyxl
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.customer import Customer
from app.models.vehicle import VehicleDetails
from app.models.transaction import Transaction
from app.models.report import PospInvoice, Target
from app.models.account import Account
from tests.integration.test_phase7_policy_booking_api import create_user_with_role
from tests.integration.test_phase12_search import make_synthetic_tx


@pytest.fixture(autouse=True)
async def cleanup_phase14_e2e(db_session: AsyncSession):
    await db_session.execute(text("DELETE FROM tbl_target;"))
    await db_session.execute(text("DELETE FROM tbl_posp_invoice;"))
    await db_session.execute(text("DELETE FROM tbl_account;"))
    await db_session.execute(text("DELETE FROM tbl_agentcommissionpayment;"))
    await db_session.execute(text("DELETE FROM tbl_transaction;"))
    await db_session.execute(text("DELETE FROM tbl_vehicledetails;"))
    await db_session.execute(text("DELETE FROM tbl_customer;"))
    await db_session.commit()
    yield
    await db_session.execute(text("DELETE FROM tbl_target;"))
    await db_session.execute(text("DELETE FROM tbl_posp_invoice;"))
    await db_session.execute(text("DELETE FROM tbl_account;"))
    await db_session.execute(text("DELETE FROM tbl_agentcommissionpayment;"))
    await db_session.execute(text("DELETE FROM tbl_transaction;"))
    await db_session.execute(text("DELETE FROM tbl_vehicledetails;"))
    await db_session.execute(text("DELETE FROM tbl_customer;"))
    await db_session.commit()


@pytest.mark.asyncio
async def test_phase14_reporting_full_e2e_lifecycle(async_client: AsyncClient, db_session: AsyncSession):
    # 1. Provision Users: Admin & Registered Agent
    admin_user, admin_headers = await create_user_with_role(db_session, "ADMIN")
    agent_user, agent_headers = await create_user_with_role(db_session, "AGENT", partner_user_id=888)
    agent_user.partner_user_id = "GST-27ABCDE1234F1Z5"
    await db_session.commit()

    # 2. Book Synthetic Policy
    cust = Customer(CustFName="RAJESH", CustLName="KHANNA", MoblieNo1="9823099999", PerAddrLine2="PUNE")
    db_session.add(cust)
    await db_session.flush()

    veh = VehicleDetails(RegistrationNo="MH12E2E001", Make_ID=1, Model_ID=1, Variant_ID=1, FinancialYear="2025-2026")
    db_session.add(veh)
    await db_session.flush()

    tx = make_synthetic_tx(
        policy_no="POL-E2E-2026-001",
        cust_veh_id=veh.CustVehId,
        customer_id=cust.CustomerId,
        gross_premium=Decimal("17700.00"),
        net_premium=Decimal("15000.00"),
        od_premium=Decimal("10000.00"),
        financial_year="2025-2026",
        agent_id=agent_user.UserId,
        branch_id=101,
    )
    tx.TransDate = datetime(2025, 4, 20, 11, 0, 0)
    tx.RiskStartdate = datetime(2025, 4, 20, 11, 0, 0)
    tx.AgentCommAmt = Decimal("1500.00")
    tx.TdsAmt = Decimal("75.00")
    tx.NetCommission = Decimal("1425.00")
    db_session.add(tx)
    await db_session.flush()

    # 3. Record Voucher & TDS in tbl_account
    acc_entry = Account(
        TransId=tx.TransanctionId,
        TransactionId=tx.TransanctionId,
        Doc_No=2025001,
        AccountDate=datetime(2025, 4, 25, 12, 0, 0),
        PaymentType="PAYMENT",
        amount=Decimal("1500.00"),
        Narration="Commission for POL-E2E-2026-001",
        LedgerMId=2113,  # Statutory TDS ledger
        BranchId=101,
        IsNill=0,
        CustVehId=veh.CustVehId,
        MonthId=4,
        EndorsementId=0,
        CreatedDate=datetime(2025, 4, 25, 12, 0, 0),
        CreatedUser=admin_user.UserName,
    )
    db_session.add(acc_entry)
    await db_session.commit()

    # 4. Verify Admin Dashboard incorporates the transaction
    resp_dash = await async_client.get(
        "/api/v1/dashboards/admin?financial_year=2025-2026&date_mode=T_Date",
        headers=admin_headers,
    )
    assert resp_dash.status_code == 200, resp_dash.text
    dash_data = resp_dash.json()
    assert dash_data["financial_year"] == "2025-2026"
    assert dash_data["totals"]["total_policies"] >= 1
    assert Decimal(str(dash_data["totals"]["total_net_premium"])) >= Decimal("15000.00")

    # 5. Generate POSP Invoice (Registered Agent: 18% GST added, 5% TDS deducted)
    inv_req = {
        "agent_id": agent_user.UserId,
        "posp_type": "POSP",
        "financial_year": "2025-2026",
        "month": "APRIL",
        "amount": "1500.00",
        "custom_invoice_no": "INV/E2E/2025-26/001",
        "is_gst_registered": True,
    }
    resp_inv = await async_client.post(
        "/api/v1/reports/posp-invoices/generate",
        json=inv_req,
        headers=admin_headers,
    )
    assert resp_inv.status_code == 201, resp_inv.text
    inv_data = resp_inv.json()

    # Math parity verification:
    # Amount = 1500.00
    # GST (18%) = 270.00
    # Grand Total = 1770.00
    # TDS (5% on 1500.00) = 75.00
    # Net Payable = 1770.00 - 75.00 = 1695.00
    assert Decimal(str(inv_data["amount"])) == Decimal("1500.00")
    assert Decimal(str(inv_data["gst_amt"])) == Decimal("270.00")
    assert Decimal(str(inv_data["grand_total"])) == Decimal("1770.00")
    assert Decimal(str(inv_data["tds_amount"])) == Decimal("75.00")
    assert Decimal(str(inv_data["net_payable"])) == Decimal("1695.00")
    assert "Rupees One Thousand Six Hundred Ninety Five Only" in inv_data["amount_in_words"]

    # 6. Verify PDF Rendering and Binary Integrity
    resp_pdf = await async_client.get(
        f"/api/v1/reports/posp-invoices/{inv_data['invoice_id']}/pdf",
        headers=admin_headers,
    )
    assert resp_pdf.status_code == 200
    assert resp_pdf.content.startswith(b"%PDF")
    assert len(resp_pdf.content) > 1000

    # 7. Verify TDS Register reports the voucher
    resp_tds = await async_client.get(
        "/api/v1/reports/statutory/tds?from_date=2025-04-01&to_date=2025-04-30",
        headers=admin_headers,
    )
    assert resp_tds.status_code == 200, resp_tds.text
    tds_data = resp_tds.json()
    assert len(tds_data["items"]) >= 1
    assert Decimal(str(tds_data["total_gross_commission"])) >= Decimal("1500.00")

    # 8. Verify Payment Advice
    resp_advice = await async_client.get(
        f"/api/v1/reports/accounting/payment-advice?financial_year=2025-2026&month=APRIL&agent_id={agent_user.UserId}",
        headers=admin_headers,
    )
    assert resp_advice.status_code == 200, resp_advice.text
    advice_data = resp_advice.json()
    assert advice_data["agent_id"] == agent_user.UserId
    assert advice_data["month"] == "APRIL"

    # 9. Verify MIS Export in XLSX format
    resp_export = await async_client.get(
        "/api/v1/reports/mis/transactions/export?format=xlsx&from_date=2025-04-01&to_date=2025-04-30",
        headers=admin_headers,
    )
    assert resp_export.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp_export.content))
    ws = wb.active
    # Header check
    assert ws.cell(row=1, column=1).value == "Trans ID"
    assert ws.cell(row=1, column=2).value == "Policy No"
    # Find our policy in rows
    found_policy = False
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == "POL-E2E-2026-001":
            found_policy = True
            break
    assert found_policy is True, "Exported XLSX did not contain POL-E2E-2026-001"

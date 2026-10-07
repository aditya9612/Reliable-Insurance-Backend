"""
Export and Document Generation Service for Phase 14.
Provides:
1. Indian Currency Number-to-Words Conversion (BLL_GetNumberIntoWord parity)
2. OpenPyXL Native Excel Spreadsheet Generation (#48D1CC headers, Century Gothic styling)
3. High-Speed Bounded-Memory CSV Streaming
4. ReportLab PDF Generation for POSP Invoices and Payment Vouchers
"""
import io
import csv
from decimal import Decimal, ROUND_HALF_UP
from typing import List, Any, Dict, Generator, Optional
from datetime import date, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


# ============================================================================
# 1. Indian Currency Number-to-Words Algorithm
# ============================================================================

_ONES = [
    "", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
    "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
    "Seventeen", "Eighteen", "Nineteen"
]

_TENS = [
    "", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"
]


def _convert_two_digits(n: int) -> str:
    if n < 20:
        return _ONES[n]
    tens = _TENS[n // 10]
    ones = _ONES[n % 10]
    return f"{tens} {ones}".strip()


def _convert_three_digits(n: int) -> str:
    hundred = n // 100
    rem = n % 100
    res = []
    if hundred > 0:
        res.append(f"{_ONES[hundred]} Hundred")
    if rem > 0:
        res.append(_convert_two_digits(rem))
    return " ".join(res).strip()


def _convert_integer_to_words(n: int) -> str:
    if n == 0:
        return "Zero"

    crore = n // 10000000
    n %= 10000000

    lakh = n // 100000
    n %= 100000

    thousand = n // 1000
    n %= 1000

    hundreds = n

    parts = []
    if crore > 0:
        parts.append(f"{_convert_integer_to_words(crore)} Crore")
    if lakh > 0:
        parts.append(f"{_convert_two_digits(lakh)} Lakh")
    if thousand > 0:
        parts.append(f"{_convert_two_digits(thousand)} Thousand")
    if hundreds > 0:
        parts.append(_convert_three_digits(hundreds))

    return " ".join(parts).strip()


def convert_number_to_words_inr(amount: Any) -> str:
    """
    Converts a numeric amount into Indian Currency Words.
    Format:
      "Rupees [Words] Only" or "Rupees [Words] and [Words] Paise Only"
    Parity with legacy BLL_GetNumberIntoWord(decimal num).
    """
    if amount is None:
        return "Rupees Zero Only"

    dec_val = Decimal(str(amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if dec_val < Decimal("0.00"):
        return f"Minus {convert_number_to_words_inr(abs(dec_val))}"

    integer_part = int(dec_val)
    paise_part = int((dec_val - Decimal(integer_part)) * 100)

    int_words = _convert_integer_to_words(integer_part)

    if paise_part > 0:
        paise_words = _convert_two_digits(paise_part)
        return f"Rupees {int_words} and {paise_words} Paise Only"
    else:
        return f"Rupees {int_words} Only"


# ============================================================================
# 2. Native OpenPyXL Spreadsheet Generator
# ============================================================================

def generate_excel_spreadsheet(
    columns: List[str],
    rows: List[List[Any]],
    sheet_name: str = "Report"
) -> bytes:
    """
    Generates a true ISO/IEC 29500 OpenXML (.xlsx) workbook in-memory.
    Preserves verified legacy styling:
    - Header background #48D1CC (Medium Turquoise)
    - Century Gothic / Arial font, bold black headers
    - Alternating row styling (#FFFFFF / #F8F9FA)
    - Auto column widths and frozen top row
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel title max 31 chars

    # Styling elements
    header_fill = PatternFill(start_color="48D1CC", end_color="48D1CC", fill_type="solid")
    header_font = Font(name="Century Gothic", size=10, bold=True, color="000000")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    regular_font = Font(name="Century Gothic", size=10, color="000000")

    thin_border_side = Side(border_style="thin", color="CCCCCC")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Write headers
    ws.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Write data rows
    for row_idx, row_data in enumerate(rows, start=2):
        formatted_row = []
        for val in row_data:
            if isinstance(val, (datetime, date)):
                formatted_row.append(val.strftime("%Y-%m-%d"))
            elif isinstance(val, Decimal):
                formatted_row.append(float(val))
            elif val is None:
                formatted_row.append("")
            else:
                formatted_row.append(val)
        ws.append(formatted_row)

        is_alt = (row_idx % 2 == 0)
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border
            if is_alt:
                cell.fill = alt_fill

            # Format numbers and currencies
            if isinstance(val, (Decimal, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(val, int):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(val, (datetime, date)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row_idx].height = 20

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Freeze header row
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ============================================================================
# 3. High-Speed Bounded-Memory CSV Streaming
# ============================================================================

def generate_csv_stream(
    columns: List[str],
    rows: Any
) -> Generator[str, None, None]:
    """
    Yields chunks of UTF-8 CSV string lines directly for StreamingResponse.
    Provides O(1) bounded server memory consumption.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Header
    writer.writerow(columns)
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    # Rows
    for row in rows:
        formatted_row = []
        for val in row:
            if isinstance(val, (datetime, date)):
                formatted_row.append(val.strftime("%Y-%m-%d"))
            elif isinstance(val, Decimal):
                formatted_row.append(f"{val:.2f}")
            elif val is None:
                formatted_row.append("")
            else:
                formatted_row.append(str(val))
        writer.writerow(formatted_row)
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)


# ============================================================================
# 4. ReportLab PDF Generation (POSP Invoices & Payment Vouchers)
# ============================================================================

def generate_posp_invoice_pdf(invoice: Dict[str, Any]) -> bytes:
    """
    Generates an official A4 POSP Payout Invoice PDF matching legacy iTextSharp layout.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "InvoiceTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        alignment=1,  # Center
        textColor=colors.HexColor("#1A365D"),
    )
    subtitle_style = ParagraphStyle(
        "InvoiceSubTitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=13,
        alignment=1,
        textColor=colors.HexColor("#4A5568"),
    )
    body_style = ParagraphStyle(
        "InvoiceBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#2D3748"),
    )
    bold_style = ParagraphStyle(
        "InvoiceBold",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#2D3748"),
    )

    story = []

    # 1. Header & Title
    story.append(Paragraph("<b>RELIABLE INSURANCE BROKING SERVICES</b>", title_style))
    story.append(Paragraph("Official POSP Agent Commission Payout Invoice", subtitle_style))
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#48D1CC"), spaceBefore=4, spaceAfter=12))

    # 2. Metadata Grid (Invoice No, Date, Agent, FY)
    meta_data = [
        [
            Paragraph(f"<b>Invoice No:</b> {invoice.get('invoice_no', '')}", body_style),
            Paragraph(f"<b>Invoice Date:</b> {invoice.get('invoice_date', '')}", body_style),
        ],
        [
            Paragraph(f"<b>POSP / Agent Name:</b> {invoice.get('agent_name', '')}", body_style),
            Paragraph(f"<b>Financial Year:</b> {invoice.get('financial_year', '')}", body_style),
        ],
        [
            Paragraph(f"<b>Agent ID:</b> {invoice.get('agent_id', '')}", body_style),
            Paragraph(f"<b>Month:</b> {invoice.get('month', '')}", body_style),
        ],
        [
            Paragraph(f"<b>POSP Type:</b> {invoice.get('posp_type', '')}", body_style),
            Paragraph(f"<b>Status:</b> {invoice.get('status', 'GENERATED')}", bold_style),
        ],
    ]
    meta_table = Table(meta_data, colWidths=[270, 265])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8F9FA")),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 16))

    # 3. Payout Line Items Table
    amount = Decimal(str(invoice.get("amount", "0.00")))
    gst_amt = Decimal(str(invoice.get("gst_amt", "0.00")))
    grand_total = Decimal(str(invoice.get("grand_total", "0.00")))
    tds_amount = Decimal(str(invoice.get("tds_amount", "0.00")))
    net_payable = Decimal(str(invoice.get("net_payable", "0.00")))

    items_data = [
        ["Sl No.", "Description / Particulars", "Tax / Rate", "Amount (INR)"],
        ["1", f"Earned POSP Commission for {invoice.get('month', '')} {invoice.get('financial_year', '')}", "Base Taxable", f"{amount:,.2f}"],
        ["2", "Goods and Services Tax (GST / IGST)", "18% (if Reg.)" if gst_amt > 0 else "0.00% (Unregistered)", f"{gst_amt:,.2f}"],
        ["", "Gross Total Invoice Value", "", f"{grand_total:,.2f}"],
        ["3", "Statutory TDS Deduction (Income Tax Sec 194H)", "5.00%", f"- {tds_amount:,.2f}"],
        ["", "Net Disbursement Payable", "", f"{net_payable:,.2f}"],
    ]

    items_table = Table(items_data, colWidths=[45, 270, 110, 110])
    items_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#48D1CC")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("ALIGN", (3, 1), (3, -1), "RIGHT"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#48D1CC")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("FONTNAME", (1, 3), (-1, 3), "Helvetica-Bold"),
        ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#EDF2F7")),
        ("FONTNAME", (1, 5), (-1, 5), "Helvetica-Bold"),
        ("BACKGROUND", (0, 5), (-1, 5), colors.HexColor("#E2E8F0")),
    ]))
    story.append(items_table)
    story.append(Spacer(1, 14))

    # 4. Amount in Words
    words = invoice.get("amount_in_words") or convert_number_to_words_inr(net_payable)
    words_para = Paragraph(f"<b>Amount in Words:</b> <i>{words}</i>", body_style)
    story.append(words_para)
    story.append(Spacer(1, 35))

    # 5. Authorization & Signature Block
    sig_data = [
        [
            Paragraph("<b>Prepared By:</b> " + invoice.get("created_by", "System"), body_style),
            Paragraph("<b>For Reliable Insurance Broking Services</b>", body_style),
        ],
        [
            "",
            Paragraph("<br/><br/><b>Authorized Signatory</b>", body_style),
        ]
    ]
    sig_table = Table(sig_data, colWidths=[270, 265])
    sig_table.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]))
    story.append(sig_table)

    doc.build(story)
    return buffer.getvalue()


def generate_voucher_pdf(voucher: Dict[str, Any]) -> bytes:
    """
    Generates an official Payment/Receipt Voucher PDF.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "VoucherTitle",
        parent=styles["Heading1"],
        fontSize=16,
        leading=20,
        alignment=1,
        textColor=colors.HexColor("#1A365D"),
    )
    body_style = ParagraphStyle(
        "VoucherBody",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#2D3748"),
    )

    story = []
    v_type = voucher.get("voucher_type", "PAYMENT VOUCHER")
    story.append(Paragraph(f"<b>RELIABLE INSURANCE — {v_type.upper()}</b>", title_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#48D1CC"), spaceBefore=2, spaceAfter=10))

    meta_data = [
        [
            Paragraph(f"<b>Voucher No:</b> {voucher.get('voucher_no', '')}", body_style),
            Paragraph(f"<b>Date:</b> {voucher.get('voucher_date', '')}", body_style),
        ],
        [
            Paragraph(f"<b>Payee / Particulars:</b> {voucher.get('payee_name', '')}", body_style),
            Paragraph(f"<b>Payment Mode:</b> {voucher.get('payment_mode', 'CASH')}", body_style),
        ],
        [
            Paragraph(f"<b>Narration:</b> {voucher.get('narration', '')}", body_style),
            Paragraph(f"<b>Created By:</b> {voucher.get('created_by', '')}", body_style),
        ],
    ]
    meta_table = Table(meta_data, colWidths=[270, 265])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8F9FA")),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 14))

    amt = Decimal(str(voucher.get("amount", "0.00")))
    amt_words = voucher.get("amount_in_words") or convert_number_to_words_inr(amt)

    val_data = [
        ["Total Voucher Amount:", f"INR {amt:,.2f}"],
        ["Amount in Words:", amt_words],
    ]
    val_table = Table(val_data, colWidths=[150, 385])
    val_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EDF2F7")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ]))
    story.append(val_table)
    story.append(Spacer(1, 40))

    sig_data = [
        [
            Paragraph("<b>Passed By / Cashier</b>", body_style),
            Paragraph("<b>Receiver's Signature</b>", body_style),
        ]
    ]
    sig_table = Table(sig_data, colWidths=[270, 265])
    sig_table.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]))
    story.append(sig_table)

    doc.build(story)
    return buffer.getvalue()
